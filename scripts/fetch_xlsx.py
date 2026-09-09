# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "google-api-python-client",
#   "google-auth",
#   "openpyxl",
# ]
# ///
"""Download an .xlsx stored on Drive (service account) and dump each sheet as JSON rows.

Dùng cho các file checklist/quan điểm được upload lên Drive dưới dạng Office file —
Sheets API KHÔNG đọc được loại này ("The document must not be an Office file").

Usage:
  uv run scripts/fetch_xlsx.py <drive_file_id> <out.xlsx>
  → ghi ra <out.xlsx> + <out.json> (mỗi sheet = list các row)

Yêu cầu: file đã share (viewer) cho service account trong credentials/google-service-account.json.
"""
import sys, io, json
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
import openpyxl

CREDS = Path(__file__).parent.parent / "credentials" / "google-service-account.json"
SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

file_id = sys.argv[1]
out_xlsx = Path(sys.argv[2])

creds = service_account.Credentials.from_service_account_file(str(CREDS), scopes=SCOPES)
drive = build("drive", "v3", credentials=creds)

meta = drive.files().get(fileId=file_id, fields="name,mimeType,size", supportsAllDrives=True).execute()
print(f"=== META: {meta}", file=sys.stderr)

buf = io.BytesIO()
req = drive.files().get_media(fileId=file_id, supportsAllDrives=True)
dl = MediaIoBaseDownload(buf, req)
done = False
while not done:
    _, done = dl.next_chunk()
out_xlsx.write_bytes(buf.getvalue())
print(f"=== saved {out_xlsx} ({out_xlsx.stat().st_size} bytes)", file=sys.stderr)

wb = openpyxl.load_workbook(out_xlsx, data_only=True)
print(f"=== SHEETS: {wb.sheetnames}", file=sys.stderr)
result = {}
for name in wb.sheetnames:
    ws = wb[name]
    rows = []
    for r in ws.iter_rows(values_only=True):
        vals = ["" if c is None else str(c).strip() for c in r]
        while vals and vals[-1] == "":
            vals.pop()
        rows.append(vals)
    while rows and not rows[-1]:
        rows.pop()
    result[name] = rows
out_json = out_xlsx.with_suffix(".json")
out_json.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"=== json -> {out_json}", file=sys.stderr)
for n, rows in result.items():
    print(f"{n}: {len(rows)} rows", file=sys.stderr)
